"""Formato visual auditable para Excel: marcar celdas estimadas vs observadas.

El patrón viene de v1: si una celda de `Monto S/` se rellenó con un
estimado TC (no con un valor observado), se pinta la fuente en rojo
para que cualquier persona que abre el Excel sepa a primera vista qué
dato es observado y cuál es imputado.

Esto es **clave para auditoría**. En reportes internos del BCRP no se
acepta "el Excel dice 2,000 soles" sin saber si lo midió el scraper o
si lo calculó un TC. Las columnas en rojo lo declaran.
"""

from __future__ import annotations

import logging
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.styles import Font

logger = logging.getLogger("scraping")

COLOR_ESTIMADO_DEFAULT = "C00000"  # rojo BCRP-style


def marcar_columnas_estimadas_excel(
    ruta_archivo: str,
    marcas_por_hoja: Mapping[str, Mapping[int, set[str]]],
    color: str = COLOR_ESTIMADO_DEFAULT,
) -> None:
    """Aplica fuente roja a celdas (fila, columna) marcadas como estimadas.

    Parámetros:
        ruta_archivo: ruta al .xlsx existente.
        marcas_por_hoja: dict {nombre_hoja: {df_idx: {col_name, ...}}}
            donde `df_idx` es el índice cero-based del DataFrame (la
            fila Excel es df_idx + 2 porque la fila 1 son headers).
        color: hex color (sin '#'). Default rojo BCRP.

    Si una hoja o columna marcada no existe en el Excel se ignora
    silenciosamente (idempotente, no falla por desfase de columnas).
    """
    wb = load_workbook(ruta_archivo)
    font = Font(color=color)

    for nombre_hoja, marcas in marcas_por_hoja.items():
        if nombre_hoja not in wb.sheetnames:
            continue
        ws = wb[nombre_hoja]
        columnas_por_nombre: dict[str, int] = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            if cell.value:
                columnas_por_nombre[str(cell.value)] = col_idx

        for df_idx, columnas in marcas.items():
            excel_row = int(df_idx) + 2
            for col_name in columnas:
                col_idx = columnas_por_nombre.get(col_name)
                if col_idx:
                    ws.cell(row=excel_row, column=col_idx).font = font

    wb.save(ruta_archivo)
