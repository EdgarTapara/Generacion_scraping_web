"""Detección determinista de potenciales duplicados (cross-source) como SEÑAL.

Este módulo **no decide ni elimina**: solo SEÑALA candidatos a duplicado (el
mismo anuncio publicado en varias fuentes/portales) para que un revisor humano
los procese. Es complementario —no sustituto— de la deduplicación exacta por
enlace que hace `core.historial`:

- `core.historial` deduplica por `enlace` canónico (exacto, automático).
- Este módulo agrupa por SIMILITUD entre campos clave (sin IA). La similitud
  NUNCA borra automáticamente: marca un grupo (`dup_grupo`) y pinta las filas
  candidatas para que el técnico decida (mantener / fusionar / eliminar).

La detección es por similitud de texto + coincidencia opcional de un valor
numérico (monto, salario), dentro de un mismo bloque de entidad y región. Las
filas conectadas se agrupan por componentes conexas (union-find) y solo los
grupos de tamaño ≥ 2 reciben id.
"""

from __future__ import annotations

import re
from collections import Counter
from difflib import SequenceMatcher
from unicodedata import normalize

import pandas as pd

# Verde claro estándar de Excel ("Good"); legible con texto negro.
VERDE_DUPLICADO = "C6EFCE"


def _sin_tildes(txt: str) -> str:
    return "".join(c for c in normalize("NFD", txt) if not re.match(r"[̀-ͯ]", c))


def _norm(valor) -> str:
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    t = _sin_tildes(str(valor)).lower()
    t = re.sub(r"[^a-z0-9 ]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _token_sort(texto: str) -> str:
    """Ordena tokens para comparar sin importar el orden de palabras."""
    return " ".join(sorted(texto.split()))


def _clave_bloque(entidad_norm: str) -> str:
    """Token más discriminante de la entidad (el más largo) para bloquear pares."""
    toks = [t for t in entidad_norm.split() if len(t) >= 3] or entidad_norm.split()
    return max(toks, key=len) if toks else ""


def _ratio(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()


def _contencion_tokens(a: str, b: str) -> float:
    """Fracción de tokens del título más corto contenidos en el otro.

    Más robusta que SequenceMatcher cuando los títulos difieren mucho en largo.
    Se usa solo como refuerzo cuando el valor numérico coincide.
    """
    ta, tb = set(a.split()), set(b.split())
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / min(len(ta), len(tb))


def _valores_coinciden(s1, s2, tol: float) -> bool:
    try:
        v1, v2 = float(s1), float(s2)
    except (TypeError, ValueError):
        return False
    if v1 <= 0 or v2 <= 0:
        return False
    return abs(v1 - v2) <= tol * max(v1, v2)


class _UnionFind:
    def __init__(self, n: int):
        self.padre = list(range(n))

    def find(self, x: int) -> int:
        while self.padre[x] != x:
            self.padre[x] = self.padre[self.padre[x]]
            x = self.padre[x]
        return x

    def union(self, a: int, b: int) -> None:
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.padre[max(ra, rb)] = min(ra, rb)


def detectar_duplicados(
    df: pd.DataFrame,
    *,
    campo_entidad: str,
    campo_titulo: str,
    campo_region: str | None = None,
    campo_valor: str | None = None,
    umbral_entidad: float = 0.80,
    umbral_titulo: float = 0.82,
    umbral_contencion_con_valor: float = 0.60,
    tol_valor: float = 0.05,
) -> list[int | None]:
    """Devuelve una lista paralela a `df`: id de grupo (1..k) o None por fila.

    Solo las filas que pertenecen a un grupo de ≥ 2 candidatos reciben id.

    Args:
        campo_entidad: columna que ancla el bloque (empresa, anunciante…).
        campo_titulo: columna de texto principal a comparar (puesto, título…).
        campo_region: si se da, solo se comparan filas de la misma región/zona.
        campo_valor: si se da, un valor numérico coincidente (monto, salario)
            refuerza el match aunque el título sea algo menos parecido.
        umbral_entidad / umbral_titulo: similitudes mínimas (0..1).
        umbral_contencion_con_valor: contención mínima de tokens del título
            cuando el match se apoya en el valor numérico.
        tol_valor: tolerancia relativa para considerar dos valores iguales.
    """
    n = len(df)
    if n == 0:
        return []

    def _col(nombre):
        return list(df[nombre]) if nombre and nombre in df.columns else [None] * n

    titulos = [_token_sort(_norm(x)) for x in _col(campo_titulo)]
    entidades = [_norm(x) for x in _col(campo_entidad)]
    entidades_sort = [_token_sort(e) for e in entidades]
    regiones = [_norm(x) for x in _col(campo_region)] if campo_region else [""] * n
    valores = _col(campo_valor)

    bloques: dict[tuple[str, str], list[int]] = {}
    for i in range(n):
        if not entidades[i] or not titulos[i]:
            continue
        clave = (regiones[i], _clave_bloque(entidades[i]))
        bloques.setdefault(clave, []).append(i)

    uf = _UnionFind(n)
    for idxs in bloques.values():
        if len(idxs) < 2:
            continue
        for a in range(len(idxs)):
            for b in range(a + 1, len(idxs)):
                i, j = idxs[a], idxs[b]
                if _ratio(entidades_sort[i], entidades_sort[j]) < umbral_entidad:
                    continue
                tit_sim = _ratio(titulos[i], titulos[j])
                if tit_sim >= umbral_titulo:
                    uf.union(i, j)
                elif campo_valor and _valores_coinciden(valores[i], valores[j], tol_valor) and (
                    _contencion_tokens(titulos[i], titulos[j]) >= umbral_contencion_con_valor
                ):
                    uf.union(i, j)

    tam = Counter(uf.find(i) for i in range(n))
    grupos: list[int | None] = [None] * n
    asignado: dict[int, int] = {}
    siguiente = 1
    for i in range(n):
        raiz = uf.find(i)
        if tam[raiz] >= 2:
            if raiz not in asignado:
                asignado[raiz] = siguiente
                siguiente += 1
            grupos[i] = asignado[raiz]
    return grupos


def aplicar_resaltado_duplicados(ws, grupos: list[int | None], color: str = VERDE_DUPLICADO) -> int:
    """Pinta las filas de datos marcadas como potencial duplicado.

    `grupos` es la lista que devuelve `detectar_duplicados`, paralela a las
    filas de datos (la fila 1 de la hoja es el encabezado). Devuelve cuántas
    filas se resaltaron. El color es solo ayuda visual: la verdad revisable son
    las columnas `dup_grupo`/`dup_score`.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    relleno = PatternFill("solid", fgColor=color)
    fuente = Font(color="000000")
    resaltadas = 0
    for offset, grupo in enumerate(grupos):
        if grupo is None:
            continue
        fila = offset + 2  # +1 encabezado, +1 base-1
        for col in range(1, ws.max_column + 1):
            celda = ws.cell(row=fila, column=col)
            celda.fill = relleno
            celda.font = fuente
            celda.alignment = Alignment(vertical="center")
        resaltadas += 1
    return resaltadas


__all__ = ["VERDE_DUPLICADO", "detectar_duplicados", "aplicar_resaltado_duplicados"]
