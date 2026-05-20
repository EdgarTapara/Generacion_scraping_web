"""Pulido visual de Excels generados (anchos de columna, header, freeze).

Estos helpers son opcionales — el sector los llama después de
`ExcelAcumulativo.escribir()` para dejar el archivo "presentable" sin
tener que repetir el mismo código en cada sector.

Convención BCRP heredada de v1:
    * Header con fondo azul oscuro (#1F4E79) + fuente blanca negrita.
    * Freeze pane en A2 (mantener cabecera al hacer scroll).
    * Columnas estimadas en rojo (#C00000) — ver core.tipo_cambio.formato.
"""

from __future__ import annotations

from typing import Mapping

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL_DEFAULT = "1F4E79"
HEADER_FONT_COLOR_DEFAULT = "FFFFFF"


def aplicar_formato_hojas(
    ruta_archivo: str,
    hojas: list[str],
    anchos_columna: Mapping[str, int] | None = None,
    ancho_default: int = 14,
    header_fill: str = HEADER_FILL_DEFAULT,
    header_font_color: str = HEADER_FONT_COLOR_DEFAULT,
    freeze_a2: bool = True,
) -> None:
    """Aplica anchos, header coloreado y freeze pane a las hojas indicadas.

    `anchos_columna` mapea nombre de columna → ancho. Columnas no listadas
    reciben `ancho_default`. Idempotente: se puede llamar varias veces.
    """
    wb = load_workbook(ruta_archivo)
    anchos = dict(anchos_columna or {})
    fill = PatternFill("solid", start_color=header_fill, end_color=header_fill)
    font = Font(bold=True, color=header_font_color)

    for nombre_hoja in hojas:
        if nombre_hoja not in wb.sheetnames:
            continue
        ws = wb[nombre_hoja]
        for col_idx, cell in enumerate(ws[1], start=1):
            col_name = str(cell.value or "")
            cell.font = font
            cell.fill = fill
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            ancho = anchos.get(col_name, ancho_default)
            ws.column_dimensions[get_column_letter(col_idx)].width = ancho
        if freeze_a2:
            ws.freeze_panes = "A2"

    wb.save(ruta_archivo)
